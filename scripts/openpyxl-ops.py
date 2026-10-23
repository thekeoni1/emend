#!/usr/bin/env python3
"""The openpyxl side of the comparison (specs/0037-measurement-and-close.md,
"The operations").

usage: python3 -I -B openpyxl-ops.py <op> <in> <out>

Runs inside the openpyxl image (tools/openpyxl-image/Containerfile), started
by scripts/compare-run.sh. Each op loads <in> with openpyxl's ordinary
load_workbook (never read_only or write_only), makes the one request that
matches an emend-copy sample, and saves to <out>. W is the first worksheet of
kind Worksheet (wb.worksheets[0]) and M its last used row (ws.max_row); every
addition goes below M, as the emend samples place theirs.

The ops and every constant they write:

  roundtrip    load; save. (emend: emend-copy <in> <out>)
  read         load; every cell of every worksheet visited by iter_rows and
               counted, the count printed as "cells=<n>"; save, since openpyxl
               has no other output to hand the harness. (--read-sample)
  edit         EDIT_NUMBER 42.5 in (M+1, A), EDIT_TEXT "emend compare" in
               (M+1, B), EDIT_BOOLEAN True in (M+1, C). (--edit-sample)
  format       FORMAT_NUMBER 1234.5 in (M+1, A), its number_format FORMAT_CODE
               "#,##0.00". (--format-sample)
  insert-rows  insert_rows(INSERT_AT 1, INSERT_COUNT 1) on W.
               (--insert-rows-sample)
  names        one defined name NAME "emend_compare" over W's $A$1.
               (--names-sample)
  table        TABLE_HEADERS "Item" and "Value" in (M+2, A) and (M+2, B),
               TABLE_ROW "first" and 1 in (M+3, A) and (M+3, B), a Table
               TABLE_NAME "EmendCompare" over A(M+2):B(M+3). (--table-sample)
  comment      a Comment COMMENT_TEXT "emend compare" by COMMENT_AUTHOR
               "emend" on (M+1, A). (--comment-sample)
  image        PNG, a 2 by 2 PNG held below as 73 bytes, added at (M+2, A).
               (--image-sample)
  add-chart    a BarChart, one series over B2:B(M+1), categories A2:A(M+1),
               added at (M+2, A). (--add-chart-sample)
  copy         copy_worksheet(W). (--copy-sample)
  chartsheet   create_chartsheet(), then add_chart of a BarChart, one series
               over W's B2:B(M+1), categories A2:A(M+1), on the chartsheet,
               which openpyxl anchors in an absoluteAnchor of ext 0 by 0
               (specs/0038-chartsheets.md, "Harness runs": the openpyxl
               fixture of --charts-sample and --chartsheet-sample).
  chart-format a LineChart, one series over B2:B(M+1), categories A2:A(M+1),
               its marker FORMAT_SYMBOL "circle" of size FORMAT_MARKER_SIZE 7
               filled FORMAT_MARKER_FILL "1F77B4", its line FORMAT_LINE_WIDTH
               28575 EMU dashed FORMAT_DASH "dash", one DataPoint at
               FORMAT_POINT 1 filled FORMAT_POINT_FILL "D62728", added at
               (M+2, A); and a PieChart over the same, one DataPoint at
               FORMAT_SLICE 0 exploded FORMAT_EXPLOSION 10, added at (M+18, A)
               (specs/0039-chart-kinds-and-formatting.md, "Harness runs":
               the openpyxl fixture of --charts-sample and
               --chart-format-sample).
  chart-analytic
               a LineChart, one series over B2:B(M+1), categories A2:A(M+1),
               its series with a Trendline of type ANALYTIC_TREND "linear"
               showing its equation and ErrorBars of ANALYTIC_ERROR_TYPE
               "fixedVal" ANALYTIC_ERROR_VALUE 1.5, the chart's dropLines,
               its y axis from ANALYTIC_MIN 0 to ANALYTIC_MAX 50 with
               majorUnit ANALYTIC_MAJOR 10, its x axis's tickLblSkip
               ANALYTIC_SKIP 2, added at (M+2, A); a date column of its own
               in the column after the first sheet's last (max_column + 1),
               serial numbers from ANALYTIC_DATE_FIRST 45292 (2024-01-01)
               on, one per row in rows 2 to M+1, plain numbers with no
               number format; and a LineChart over B2:B(M+1) whose
               categories are that column, with a DateAxis as its x axis,
               its baseTimeUnit ANALYTIC_TIME_UNIT "days" and its majorUnit
               ANALYTIC_DATE_MAJOR 2 of majorTimeUnit ANALYTIC_TIME_UNIT,
               its y axis of logBase ANALYTIC_LOG_BASE 10, added at (M+18, A)
               (specs/0040-analytic-chart-elements.md, "Harness runs": the
               openpyxl fixture of --charts-sample and --chart-axis-sample).
  chart-text   a BarChart, one series over B2:B(M+1), categories A2:A(M+1),
               its style TEXT_STYLE 26, its title TEXT_TITLE "emend text",
               its x axis's txPr of sz TEXT_AXIS_SIZE 1400 bold with the latin
               typeface TEXT_FONT "Arial", and data labels showing the values
               with a txPr of sz TEXT_LABEL_SIZE 900, added at (M+2, A)
               (specs/0041-chart-text-and-edits.md, "Harness runs": the
               openpyxl fixture of --charts-sample, --chart-text-sample and
               --chart-edit-sample).
  named-style  a NamedStyle NAMED_STYLE_NAME "emend named style" with a bold
               Font and a linear GradientFill of degree NAMED_STYLE_DEGREE 90
               over the two stops NAMED_STYLE_STOPS "FF336699" and
               "FFDDEEFF", added with add_named_style and set on the cell at
               (M+1, A) through cell.style
               (specs/0042-gradient-fills-and-named-styles.md, "Harness
               runs": the openpyxl fixture of --namedstyles-sample,
               --gradient-sample and --namedstyle-sample).

What it prints, on standard output: for read, the line "cells=<n>"; for an
exception, one line holding the exception's class name only, never its
message, which may quote workbook content; and last, whatever happened after
the op began, the line "openpyxl=<version>\\tpeak_rss_kb=<ru_maxrss>" from
resource.getrusage(RUSAGE_SELF). Nothing it prints names a path or a cell
value.

Exits: 0 when the op saved its output; 1 on an exception (its class printed);
137 on MemoryError, so the runner classes it as the cap, a MemoryError raised
while the script reports another exception included (specs/0037, ruling Y:
that path writes its two lines with bytes made before the op, by os.write,
and leaves by os._exit, so it needs no new memory to end); 2, with the one line
"unknown op" and nothing else, for an op it does not know or a wrong number of
arguments. openpyxl's own warnings go to standard error, which the runner
keeps in the run directory and never prints.
"""
import os
import resource
import sys

OPS = ("roundtrip", "read", "edit", "format", "insert-rows", "names", "table",
       "comment", "image", "add-chart", "copy", "chartsheet", "chart-format",
       "chart-analytic", "chart-text", "named-style")

EDIT_NUMBER = 42.5
EDIT_TEXT = "emend compare"
EDIT_BOOLEAN = True
FORMAT_NUMBER = 1234.5
FORMAT_CODE = "#,##0.00"
INSERT_AT = 1
INSERT_COUNT = 1
NAME = "emend_compare"
TABLE_HEADERS = ("Item", "Value")
TABLE_ROW = ("first", 1)
TABLE_NAME = "EmendCompare"
COMMENT_TEXT = "emend compare"
COMMENT_AUTHOR = "emend"
FORMAT_SYMBOL = "circle"
FORMAT_MARKER_SIZE = 7
FORMAT_MARKER_FILL = "1F77B4"
FORMAT_LINE_WIDTH = 28575
FORMAT_DASH = "dash"
FORMAT_POINT = 1
FORMAT_POINT_FILL = "D62728"
FORMAT_SLICE = 0
FORMAT_EXPLOSION = 10
ANALYTIC_TREND = "linear"
ANALYTIC_ERROR_TYPE = "fixedVal"
ANALYTIC_ERROR_VALUE = 1.5
ANALYTIC_MIN = 0
ANALYTIC_MAX = 50
ANALYTIC_MAJOR = 10
ANALYTIC_SKIP = 2
ANALYTIC_TIME_UNIT = "days"
ANALYTIC_DATE_MAJOR = 2
ANALYTIC_LOG_BASE = 10
ANALYTIC_DATE_FIRST = 45292
TEXT_STYLE = 26
TEXT_TITLE = "emend text"
TEXT_AXIS_SIZE = 1400
TEXT_FONT = "Arial"
TEXT_LABEL_SIZE = 900
NAMED_STYLE_NAME = "emend named style"
NAMED_STYLE_DEGREE = 90
NAMED_STYLE_STOPS = ("FF336699", "FFDDEEFF")
# A 2 by 2 RGB PNG, 73 bytes: red and blue on each row.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000002000000020802000000fdd49a73"
    "000000104944415478da63f8cf0004ff192014001bf203fd3e975da40000000049"
    "454e44ae426082"
)


def run(op, source, target):
    import io
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.axis import ChartLines, DateAxis
    from openpyxl.chart.error_bar import ErrorBars
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.text import RichText
    from openpyxl.chart.trendline import Trendline
    from openpyxl.chart.marker import DataPoint, Marker
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.drawing.text import CharacterProperties, Font, Paragraph, ParagraphProperties
    from openpyxl.comments import Comment
    from openpyxl.styles import Font as CellFont, GradientFill, NamedStyle
    from openpyxl.drawing.image import Image
    from openpyxl.utils import quote_sheetname
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    wb = openpyxl.load_workbook(source)
    ws = wb.worksheets[0]
    m = ws.max_row
    if op == "read":
        cells = 0
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                cells += len(row)
        print("cells=%d" % cells)
    elif op == "edit":
        ws.cell(row=m + 1, column=1, value=EDIT_NUMBER)
        ws.cell(row=m + 1, column=2, value=EDIT_TEXT)
        ws.cell(row=m + 1, column=3, value=EDIT_BOOLEAN)
    elif op == "format":
        ws.cell(row=m + 1, column=1, value=FORMAT_NUMBER).number_format = FORMAT_CODE
    elif op == "insert-rows":
        ws.insert_rows(INSERT_AT, INSERT_COUNT)
    elif op == "names":
        ref = "%s!$A$1" % quote_sheetname(ws.title)
        wb.defined_names[NAME] = DefinedName(NAME, attr_text=ref)
    elif op == "table":
        ws.cell(row=m + 2, column=1, value=TABLE_HEADERS[0])
        ws.cell(row=m + 2, column=2, value=TABLE_HEADERS[1])
        ws.cell(row=m + 3, column=1, value=TABLE_ROW[0])
        ws.cell(row=m + 3, column=2, value=TABLE_ROW[1])
        ws.add_table(Table(displayName=TABLE_NAME, ref="A%d:B%d" % (m + 2, m + 3)))
    elif op == "comment":
        ws.cell(row=m + 1, column=1).comment = Comment(COMMENT_TEXT, COMMENT_AUTHOR)
    elif op == "image":
        ws.add_image(Image(io.BytesIO(PNG)), "A%d" % (m + 2))
    elif op == "add-chart":
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=m + 1), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=m + 1))
        ws.add_chart(chart, "A%d" % (m + 2))
    elif op == "copy":
        wb.copy_worksheet(ws)
    elif op == "chartsheet":
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=m + 1), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=m + 1))
        wb.create_chartsheet().add_chart(chart)
    elif op == "chart-format":
        values = Reference(ws, min_col=2, min_row=2, max_row=m + 1)
        categories = Reference(ws, min_col=1, min_row=2, max_row=m + 1)
        line = LineChart()
        line.add_data(values, titles_from_data=False)
        line.set_categories(categories)
        series = line.series[0]
        series.marker = Marker(symbol=FORMAT_SYMBOL, size=FORMAT_MARKER_SIZE)
        series.marker.graphicalProperties = GraphicalProperties(solidFill=FORMAT_MARKER_FILL)
        series.graphicalProperties = GraphicalProperties(
            ln=LineProperties(w=FORMAT_LINE_WIDTH, prstDash=FORMAT_DASH))
        point = DataPoint(idx=FORMAT_POINT)
        point.graphicalProperties = GraphicalProperties(solidFill=FORMAT_POINT_FILL)
        series.dPt.append(point)
        ws.add_chart(line, "A%d" % (m + 2))
        pie = PieChart()
        pie.add_data(values, titles_from_data=False)
        pie.set_categories(categories)
        pie.series[0].dPt.append(DataPoint(idx=FORMAT_SLICE, explosion=FORMAT_EXPLOSION))
        ws.add_chart(pie, "A%d" % (m + 18))
    elif op == "chart-analytic":
        values = Reference(ws, min_col=2, min_row=2, max_row=m + 1)
        categories = Reference(ws, min_col=1, min_row=2, max_row=m + 1)
        line = LineChart()
        line.add_data(values, titles_from_data=False)
        line.set_categories(categories)
        series = line.series[0]
        series.trendline = Trendline(trendlineType=ANALYTIC_TREND, dispEq=True)
        series.errBars = ErrorBars(errValType=ANALYTIC_ERROR_TYPE, val=ANALYTIC_ERROR_VALUE)
        line.dropLines = ChartLines()
        line.y_axis.scaling.min = ANALYTIC_MIN
        line.y_axis.scaling.max = ANALYTIC_MAX
        line.y_axis.majorUnit = ANALYTIC_MAJOR
        line.x_axis.tickLblSkip = ANALYTIC_SKIP
        ws.add_chart(line, "A%d" % (m + 2))
        date_column = ws.max_column + 1
        for row in range(2, m + 2):
            ws.cell(row=row, column=date_column, value=ANALYTIC_DATE_FIRST + row - 2)
        dates = Reference(ws, min_col=date_column, min_row=2, max_row=m + 1)
        dated = LineChart()
        dated.add_data(values, titles_from_data=False)
        dated.set_categories(dates)
        dated.y_axis.crossAx = 500
        dated.x_axis = DateAxis(crossAx=100)
        dated.x_axis.baseTimeUnit = ANALYTIC_TIME_UNIT
        dated.x_axis.majorUnit = ANALYTIC_DATE_MAJOR
        dated.x_axis.majorTimeUnit = ANALYTIC_TIME_UNIT
        dated.y_axis.scaling.logBase = ANALYTIC_LOG_BASE
        ws.add_chart(dated, "A%d" % (m + 18))
    elif op == "chart-text":
        def text_properties(properties):
            return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=properties),
                                         endParaRPr=CharacterProperties())])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=m + 1), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=m + 1))
        chart.style = TEXT_STYLE
        chart.title = TEXT_TITLE
        chart.x_axis.txPr = text_properties(
            CharacterProperties(sz=TEXT_AXIS_SIZE, b=True, latin=Font(typeface=TEXT_FONT)))
        chart.dataLabels = DataLabelList(showVal=True)
        chart.dataLabels.txPr = text_properties(CharacterProperties(sz=TEXT_LABEL_SIZE))
        ws.add_chart(chart, "A%d" % (m + 2))
    elif op == "named-style":
        style = NamedStyle(name=NAMED_STYLE_NAME)
        style.font = CellFont(bold=True)
        style.fill = GradientFill(type="linear", degree=NAMED_STYLE_DEGREE, stop=NAMED_STYLE_STOPS)
        wb.add_named_style(style)
        ws.cell(row=m + 1, column=1).style = NAMED_STYLE_NAME
    wb.save(target)
    return openpyxl.__version__


def last_line(version):
    rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    print("openpyxl=%s\tpeak_rss_kb=%d" % (version, rss))


# The cap's two lines, made before the op (ruling Y): the class name, and the
# last line's prefix, its version filled in once openpyxl is imported.
MEMORY_LINE = b"MemoryError\n"
NO_RSS = b"-\n"
LAST_PREFIX = [b"openpyxl=-\tpeak_rss_kb="]


def cap_exit():
    """Ends the script with 137 after a MemoryError, wherever it was raised:
    what print has buffered is flushed if it can be, the two lines go out by
    os.write, the peak RSS as "-" if even its digits cannot be made, and
    os._exit leaves whatever happens on the way."""
    try:
        try:
            sys.stdout.flush()
        except BaseException:
            pass
        os.write(1, MEMORY_LINE)
        os.write(1, LAST_PREFIX[0])
        try:
            rss = b"%d\n" % resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        except MemoryError:
            rss = NO_RSS
        os.write(1, rss)
    finally:
        os._exit(137)


def main(argv):
    if len(argv) != 3 or argv[0] not in OPS:
        print("unknown op")
        return 2
    version = "-"
    try:
        import openpyxl
        version = openpyxl.__version__
        LAST_PREFIX[0] = ("openpyxl=%s\tpeak_rss_kb=" % version).encode()
        run(*argv)
    except MemoryError:
        cap_exit()
    except Exception as err:
        print(type(err).__name__)
        last_line(version)
        return 1
    last_line(version)
    return 0


if __name__ == "__main__":
    # A MemoryError raised while another exception is reported, or anywhere
    # outside the op, ends as the cap too (ruling Y).
    try:
        code = main(sys.argv[1:])
    except MemoryError:
        cap_exit()
    sys.exit(code)
